# app/utils/excel.py
from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


HEADERS = [
    "ID", "Дата", "Имя", "Компания", "Кто",
    "Интерес на стенде", "Направления", "Что интересует",
    "Телефон", "Email", "После выставки",
]


def _survey_to_row(s) -> list:
    """Превращает ORM-объект анкеты в список значений для строки Excel."""
    return [
        s.id,
        s.created_at.strftime("%d.%m.%Y %H:%M") if s.created_at else "",
        s.name or "",
        s.company or "",
        s.role or "",
        ", ".join(s.stand_interest or []),
        ", ".join(s.directions or []),
        ", ".join(s.interest or []),
        s.phone or "",
        s.email or "",
        s.followup or "",
    ]


def surveys_to_xlsx(surveys: Iterable) -> BytesIO:
    """
    Собирает .xlsx со всеми анкетами и возвращает BytesIO.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Анкеты"

    # Шапка
    ws.append(HEADERS)
    for col_idx, _ in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Данные
    for s in surveys:
        ws.append(_survey_to_row(s))

    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_length = 0
        for cell in column_cells:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 50)

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer